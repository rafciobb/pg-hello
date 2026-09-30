"""Buduje plik XLSX content planu EuroFrance z danych w seo/eurofrance/data/plan_q*.py.

Uruchomienie (z katalogu repozytorium):
    python3 seo/eurofrance/tools/build_plan.py

Skrypt waliduje dane (URL-e kategorii z drzewa w prompcie, unikalność fraz i slugów,
5 wpisów na miesiąc, długość title) i zapisuje seo/eurofrance/03-content-plan-eurofrance.xlsx.
"""
import glob
import importlib.util
import os
import re
import sys
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PROMPT = os.path.join(ROOT, "01-prompt-content-plan-roczny.md")
OUT = os.path.join(ROOT, "03-content-plan-eurofrance.xlsx")

COLUMNS = [
    ("ID", "id", 12),
    ("Miesiąc i sugerowany tydzień", "termin", 16),
    ("Klaster tematyczny", "klaster", 22),
    ("Rola w klastrze", "rola", 10),
    ("Kategoria sklepu", "kategoria", 28),
    ("Tytuł H1", "h1", 42),
    ("Title tag (≤ ok. 60 zn.)", "title", 34),
    ("Proponowany slug", "slug", 30),
    ("Fraza główna", "fraza", 24),
    ("Frazy wspierające / long-tail", "wspierajace", 40),
    ("Pytania użytkowników do pokrycia", "pytania", 48),
    ("Intencja i etap lejka", "intencja", 22),
    ("Typ i format treści", "typ", 36),
    ("Encja główna + atrybuty EAV + encje powiązane", "encje", 50),
    ("Wyróżnik / information gain", "wyroznik", 40),
    ("Linki do kategorii (URL → anchor)", "linki_kat", 60),
    ("Linki wewnętrzne blog", "linki_blog", 32),
    ("CTA / most do konwersji", "cta", 34),
    ("Sezonowość i uzasadnienie terminu", "sezon", 32),
    ("Potencjał ruchu", "pr", 11),
    ("Konkurencyjność SERP", "kd", 12),
    ("Wartość biznesowa (1–5)", "wb", 10),
    ("Priorytet", None, 9),
    ("Ryzyko / ostrożność", "ryzyko", 32),
    ("Sugerowana objętość + dane strukturalne", "objetosc", 26),
    ("Komentarz SEO", "komentarz", 50),
]

FONT = "Arial"
HEADER_FILL = PatternFill("solid", start_color="1F3A5F")
MONTH_FILLS = ["FFFFFF", "F2F6FA"]
THIN = Side(style="thin", color="C9D3DD")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BLUE = Font(name=FONT, size=10, color="0000FF")
BLACK = Font(name=FONT, size=10, color="000000")
GREEN = Font(name=FONT, size=10, color="008000")
BOLD = Font(name=FONT, size=10, bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(wrap_text=True, vertical="top", horizontal="center")


def load_rows():
    rows = []
    for path in sorted(glob.glob(os.path.join(ROOT, "data", "plan_q*.py"))):
        spec = importlib.util.spec_from_file_location(os.path.basename(path)[:-3], path)
        mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mod)
        rows.extend(mod.ROWS)
    return rows


def validate(rows):
    with open(PROMPT, encoding="utf-8") as fh:
        tree = {line.strip() for line in fh if line.startswith("https://www.eurofrance.pl/autoczesci")}
    errors = []
    for r in rows:
        for url in re.findall(r"https://\S+?\.html", r["linki_kat"]):
            if url not in tree:
                errors.append(f"{r['id']}: URL spoza drzewa: {url}")
        n_links = len(re.findall(r"https://\S+?\.html", r["linki_kat"]))
        if not 2 <= n_links <= 5:
            errors.append(f"{r['id']}: {n_links} linków do kategorii (wymagane 2–5)")
        if len(r["title"]) > 62:
            errors.append(f"{r['id']}: title ma {len(r['title'])} znaków")
        if not re.search(r"EF-\d{4}-\d{2}", r["linki_blog"]):
            errors.append(f"{r['id']}: brak linku blog ↔ blog")
    for key in ("id", "fraza", "slug"):
        for val, cnt in Counter(r[key].lower() for r in rows).items():
            if cnt > 1:
                errors.append(f"Duplikat {key}: {val}")
    months = Counter(r["id"][3:7] for r in rows)
    for month, cnt in months.items():
        if cnt != 5:
            errors.append(f"Miesiąc {month}: {cnt} wpisów (wymagane 5)")
    return errors


def build(rows):
    wb = Workbook()

    # ---------------------------------------------------------------- Content plan
    ws = wb.active
    ws.title = "Content plan"
    for c, (name, _, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.row_dimensions[1].height = 42

    prio_col = [i for i, (_, key, _) in enumerate(COLUMNS, start=1) if key is None][0]
    month_index = {m: i for i, m in enumerate(dict.fromkeys(r["id"][3:7] for r in rows))}
    n = len(rows)
    for i, r in enumerate(rows, start=2):
        fill = PatternFill("solid", start_color=MONTH_FILLS[month_index[r["id"][3:7]] % 2])
        for c, (_, key, _) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=i, column=c)
            if key is None:
                cell.value = f"=INDEX(Scoring!$J$10:$J${9 + n},MATCH($A{i},Scoring!$A$10:$A${9 + n},0))"
                cell.font = GREEN
                cell.alignment = CENTER
            else:
                cell.value = r[key]
                cell.font = BLUE if key == "wb" else BLACK
                cell.alignment = CENTER if key in ("id", "rola", "pr", "kd", "wb") else WRAP
            cell.fill = fill
            cell.border = BORDER
        ws.cell(row=i, column=1).font = BOLD
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{n + 1}"

    # ---------------------------------------------------------------- Scoring
    sc = wb.create_sheet("Scoring")
    sc["A1"] = "Model priorytetyzacji (Etap A, pkt 6)"
    sc["A1"].font = Font(name=FONT, size=12, bold=True)
    params = [
        ("Waga", "Faza „Start” (XI 2026 – I 2027)", "Faza „Standard” (od II 2027)"),
        ("Wartość biznesowa (WB)", 0.4, 0.4),
        ("Potencjał ruchu (PR)", 0.2, 0.3),
        ("Łatwość rankowania (Ł)", 0.3, 0.2),
        ("Pilność sezonowa (PS)", 0.1, 0.1),
    ]
    for r_off, row in enumerate(params, start=2):
        for c_off, val in enumerate(row, start=1):
            cell = sc.cell(row=r_off, column=c_off, value=val)
            cell.border = BORDER
            if r_off == 2:
                cell.font = BOLD
            elif c_off > 1:
                cell.font = BLUE
                cell.number_format = "0%"
                cell.fill = PatternFill("solid", start_color="FFFF00")
            else:
                cell.font = BLACK
    sc["E2"], sc["F2"] = "Próg", "Wartość"
    sc["E2"].font = sc["F2"].font = BOLD
    sc["E3"], sc["F3"] = "P1 od wyniku", 3.8
    sc["E4"], sc["F4"] = "P2 od wyniku", 3.0
    for ref in ("F3", "F4"):
        sc[ref].font = BLUE
        sc[ref].fill = PatternFill("solid", start_color="FFFF00")
    for ref in ("E3", "E4"):
        sc[ref].font = BLACK
    sc["E5"] = "Pillar zawsze = P1 (rola strukturalna). Poniżej progu P2 = P3."
    sc["E5"].font = Font(name=FONT, size=9, italic=True)
    sc["A7"] = ("Oceny 1–5. PR: W=5, Ś-W=4, Ś=3, N=2, nisza=1. Ł (odwrotność konkurencyjności): N=5, Ś=3, "
                "Ś-W=2, W=1. PS: okno sezonowe=5, blisko sezonu=4, evergreen=3, poza sezonem=1. "
                "Oceny PR i Ł to szacunki bez danych z narzędzi (Senuto/Ahrefs/GSC), do weryfikacji.")
    sc["A7"].font = Font(name=FONT, size=9, italic=True)
    sc["A7"].alignment = WRAP
    sc.merge_cells("A7:J7")
    sc.row_dimensions[7].height = 40

    heads = ["ID", "Rola", "Faza", "WB", "PR", "Ł", "PS", "Wynik", "Uwagi", "Priorytet"]
    for c, h in enumerate(heads, start=1):
        cell = sc.cell(row=9, column=c, value=h)
        cell.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = BORDER
    for i, r in enumerate(rows, start=10):
        plan_row = i - 8
        phase = "Start" if r["id"][3:7] in ("2611", "2612", "2701") else "Standard"
        values = [
            (f"='Content plan'!A{plan_row}", GREEN),
            (f"='Content plan'!D{plan_row}", GREEN),
            (phase, BLUE),
            (f"='Content plan'!V{plan_row}", GREEN),
            (r["s_pr"], BLUE),
            (r["s_l"], BLUE),
            (r["s_ps"], BLUE),
            (f'=IF(C{i}="Start",D{i}*$B$3+E{i}*$B$4+F{i}*$B$5+G{i}*$B$6,'
             f'D{i}*$C$3+E{i}*$C$4+F{i}*$C$5+G{i}*$C$6)', BLACK),
            (f'=IF(AND(B{i}="Pillar",H{i}<$F$3),"pillar: P1 mimo wyniku","")', BLACK),
            (f'=IF(B{i}="Pillar","P1",IF(H{i}>=$F$3,"P1",IF(H{i}>=$F$4,"P2","P3")))', BLACK),
        ]
        for c, (val, font) in enumerate(values, start=1):
            cell = sc.cell(row=i, column=c, value=val)
            cell.font = font
            cell.border = BORDER
            cell.alignment = CENTER
        sc.cell(row=i, column=8).number_format = "0.00"
    for c, w in zip("ABCDEFGHIJ", (13, 26, 26, 8, 8, 8, 8, 9, 24, 10)):
        sc.column_dimensions[c].width = w
    sc.freeze_panes = "A10"

    # ---------------------------------------------------------------- Legenda
    lg = wb.create_sheet("Legenda")
    lines = [
        ("Content plan bloga EuroFrance: listopad 2026 – październik 2027", "title"),
        ("Stan: Etap B, część 1 (XI 2026 – I 2027). Kolejne kwartały są dopisywane do tego samego pliku.", ""),
        ("", ""),
        ("Kolory", "h"),
        ("Niebieski tekst = wartości wpisane ręcznie (oceny, wagi, wartość biznesowa); można je zmieniać.", ""),
        ("Zielony tekst = wartości pobierane z innego arkusza (formuły).", ""),
        ("Żółte tło = kluczowe założenia (wagi, progi priorytetu).", ""),
        ("", ""),
        ("Jak czytać kolumny", "h"),
        ("ID: EF-RRMM-NN (rok, miesiąc, numer w miesiącu).", ""),
        ("Rola: Pillar = kompendium klastra; Cluster = konkretny problem w klastrze; Support = wąski long-tail / szybka wygrana.", ""),
        ("Potencjał ruchu / Konkurencyjność: W / Ś-W / Ś / N / Nisza; „szac.” = szacunek bez danych z narzędzi, do weryfikacji w Senuto/Ahrefs/GSC.", ""),
        ("Priorytet: liczony w arkuszu Scoring (wagi i progi w żółtych komórkach). Pillary zawsze P1.", ""),
        ("Linki do kategorii: wyłącznie URL-e z drzewa kategorii sklepu (walidowane skryptem).", ""),
        ("Linki wewnętrzne blog: „Linkuje do” = linki wychodzące z wpisu; „Zaktualizować” = które wcześniejsze wpisy trzeba uzupełnić o link.", ""),
        ("[INTL] w komentarzu = temat uniwersalny, kandydat do tłumaczenia na rynki zagraniczne.", ""),
        ("", ""),
        ("Klastry", "h"),
        ("K1 Świadomy zakup i dobór części · K2 Rozrząd i napęd pasowy · K3 Sprzęgło, dwumasa i skrzynia biegów · K4 Hamulce · "
         "K5 Zawieszenie i układ kierowniczy · K6 Chłodzenie, ogrzewanie i klimatyzacja · K7 Elektronika samochodowa · "
         "K8 Elektryka, rozruch i ładowanie · K9 Silnik: diagnostyka, emisja i układ paliwowy · K10 Karoseria, oświetlenie i widoczność", ""),
        ("", ""),
        ("Założenia (Etap A)", "h"),
        ("Z1 Marki francuskie ok. 75–85% tematów modelowych. Z2 Brak danych sprzedażowych: wartość biznesowa szacowana. "
         "Z3 Brak danych GSC/Senuto/Ahrefs: szacunki jakościowe. Z4 Dobór po VIN, wyszukiwarka po numerze OE i gwarancja: do potwierdzenia (CTA neutralne). "
         "Z5 Ekspert po stronie klienta: opcjonalny. Z6 Części używane: margines, bez dedykowanych tematów. Z7 Liczby magazynowe i oceny: potwierdzić przed użyciem.", ""),
        ("", ""),
        ("Przed publikacją każdego wpisu", "h"),
        ("Zweryfikować w narzędziach frazę główną i long-taile; potwierdzić dane liczbowe (interwały, koszty jako widełki); "
         "sprawdzić flagę ryzyka; dodać linki zwrotne we wcześniejszych wpisach wg kolumny „Linki wewnętrzne blog”.", ""),
    ]
    for i, (text, style) in enumerate(lines, start=1):
        cell = lg.cell(row=i, column=1, value=text)
        if style == "title":
            cell.font = Font(name=FONT, size=13, bold=True)
        elif style == "h":
            cell.font = Font(name=FONT, size=11, bold=True, color="1F3A5F")
        else:
            cell.font = BLACK
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    lg.column_dimensions["A"].width = 140

    wb.move_sheet("Legenda", offset=-2)
    wb.active = 1
    wb.save(OUT)


if __name__ == "__main__":
    data = load_rows()
    problems = validate(data)
    if problems:
        print("BŁĘDY WALIDACJI:")
        print("\n".join(problems))
        sys.exit(1)
    build(data)
    print(f"OK: {len(data)} wierszy → {OUT}")
